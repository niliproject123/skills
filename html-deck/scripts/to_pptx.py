# -*- coding: utf-8 -*-
"""Export a deck folder to a .pptx that looks exactly like the HTML.

Every <section class="slide"> is shot in Chromium and placed on its own 16:9 pptx
slide. The result is pixel-identical to the deck and its text is NOT editable in
PowerPoint - that trade is the whole point. A pptx built from native shapes cannot
reproduce these layouts, and you cannot see it to check.

A slide in the browser is a white card: 1px var(--line) border, 6px corners, a soft
shadow, on #E9EDF2 paper. --frame decides how much of that the pptx keeps:

    all     (default) every slide is the card, on deck paper - what the deck
            looks like. Costs ~3% of the width to the margin.
    fitted  only slides that grew past 720px are carded; the rest go edge to
            edge. Biggest text, and the fitted ones stop looking like accidents.
    none    edge to edge everywhere, no border, no paper.

Slides that grew past 720px are fitted by height and centred, so nothing is ever
cropped; they end up smaller and the run reports each one by name. They are shot
one step sharper, since they are the ones being scaled down.

Needs a deck server already running - the decks link /_deck/deck.css, which only
resolves over http. Start one first:
    python <skill>/scripts/serve.py <deck-folder> 8899

    PYTHONIOENCODING=utf-8 python <skill>/scripts/to_pptx.py <deck-folder> out.pptx
    ... --sheet check.png      also write a contact sheet you can look at

Requires: playwright (+ chromium), python-pptx; pillow only for --sheet.
"""
import argparse
import glob
import io
import os
import re
import sys
import urllib.error
import urllib.request

# A Windows console is cp1252, and a deck title is not: without this the script dies while
# printing the title it just read off the wire. Decks are Hebrew, Arabic, Greek more often than not.
if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')
if hasattr(sys.stderr, 'reconfigure'):
    sys.stderr.reconfigure(encoding='utf-8')


EMU_IN = 914400
SLIDE_W = int(13.3333 * EMU_IN)          # 16:9, the shape of a .slide
SLIDE_H = int(7.5 * EMU_IN)

PAGE_BG = (0xE9, 0xED, 0xF2)             # deck.css html,body - the paper
TALL = 721                               # a .grow slide that outgrew --slide-h
MARGIN = 22                              # paper around a carded slide; the deck
                                         # shadow is 0 2px 14px, so 22 clears it

# The deck chrome is screen-only furniture. @media print drops most of it, these
# rules drop the rest. Paper colour stays, and .slide keeps its card styling --
# stripping that is what made earlier exports look unlike the deck.
CAPTURE_CSS = """
#deckbar,#bar,#drawer,#slidebar,.undo-btn{display:none!important}
body{padding:0!important;margin:0!important;background:#E9EDF2!important}
*{scrollbar-width:none!important}
/* fit.js zooms every .slide to fill the reader's window. Left alone it makes a
   1280x720 slide measure 1396x785 here, so every slide reads as "taller than 720"
   and gets shrunk to fit. Capture at the design size instead. */
.slide{zoom:1!important}
.pptx-frame{display:inline-block;width:max-content;background:#E9EDF2;padding:%dpx}
.pptx-frame>.slide{margin:0!important}
.pptx-flat{margin:0!important;box-shadow:none!important;border:0!important;
           border-radius:0!important}
""" % MARGIN

# Runs in the page. Cards the slides this mode wants carded, flattens the rest,
# and tags whichever element is the one to shoot so python can find it in order.
PREPARE_JS = """
(mode) => {
  const out = [];
  document.querySelectorAll('section.slide').forEach(s => {
    const tall = s.getBoundingClientRect().height > %d;
    const framed = mode === 'all' || (mode === 'fitted' && tall);
    let target = s;
    if (framed) {
      const w = document.createElement('div');
      w.className = 'pptx-frame';
      s.parentNode.insertBefore(w, s);
      w.appendChild(s);
      target = w;
    } else {
      s.classList.add('pptx-flat');
    }
    target.setAttribute('data-shot', s.dataset.no || '??');
    out.push({no: s.dataset.no || '??', framed: framed, tall: tall});
  });
  return out;
}
""" % TALL


def require_server(base, decks):
    """Fail loudly and say what to run. A silent fallback here would export a
    deck with no stylesheet, which looks like a broken pptx, not a missing one."""
    here = os.path.dirname(os.path.abspath(__file__))
    for deck in decks:
        url = base + deck
        try:
            with urllib.request.urlopen(url, timeout=5) as r:
                if r.status != 200:
                    raise SystemExit("%s -> HTTP %d" % (url, r.status))
        except urllib.error.URLError as e:
            raise SystemExit(
                "cannot reach %s (%s)\n"
                "start the deck server first:\n"
                "    python %s <deck-folder> <port>"
                % (url, e, os.path.join(here, "serve.py")))


def deck_title(base, deck):
    """The deck's own <title> becomes the presentation title, so the file shows
    its Hebrew name in PowerPoint's title bar and in Explorer's Title column."""
    with urllib.request.urlopen(base + deck, timeout=5) as r:
        html = r.read().decode("utf-8", "replace")
    m = re.search(r"<title>(.*?)</title>", html, re.S)
    return m.group(1).strip() if m else None


def open_deck(ctx, url, mode):
    page = ctx.new_page()
    page.emulate_media(media="print")
    page.goto(url, wait_until="networkidle")
    page.add_style_tag(content=CAPTURE_CSS)
    page.evaluate("document.fonts.ready")
    page.wait_for_timeout(400)          # webfont swap, after fonts.ready
    plan = page.evaluate(PREPARE_JS, mode)
    if not plan:
        raise SystemExit("no slide sections found at " + url)
    return page, plan


def capture(p, base, decks, shots_dir, scale, mode):
    """Two contexts, because device_scale_factor is fixed per context: the scaled
    down slides are shot one step sharper than the ones that fill the frame."""
    os.makedirs(shots_dir, exist_ok=True)
    browser = p.chromium.launch()
    view = {"width": 1440, "height": 900}      # wider than a 1280 card + margin,
    ctxs = {}                                  # so nothing has to scroll sideways

    def ctx_at(s):
        if s not in ctxs:
            ctxs[s] = browser.new_context(viewport=view, device_scale_factor=s)
        return ctxs[s]

    shots = []
    for deck in decks:
        stem = os.path.splitext(deck)[0]
        url = base + deck
        pages = {}
        _, plan = open_deck(ctx_at(scale), url, mode)
        pages[scale] = _
        for item in plan:
            s = scale + 1 if item["tall"] else scale
            if s not in pages:
                pages[s], _ignored = open_deck(ctx_at(s), url, mode)
            page = pages[s]
            el = page.query_selector('[data-shot="%s"]' % item["no"])
            if el is None:
                raise SystemExit("slide %s vanished from %s" % (item["no"], deck))
            box = el.bounding_box()
            path = os.path.join(shots_dir, "%s-%s.png" % (stem, item["no"]))
            el.screenshot(path=path)
            shots.append((deck, item["no"], path, box["width"], box["height"],
                          item["tall"]))
            print("  %-14s %s  %dx%d  %s @%dx"
                  % (stem, item["no"], box["width"], box["height"],
                     "card" if item["framed"] else "full", s))
        for pg in pages.values():
            pg.close()
    browser.close()
    return shots


def build(shots, out, title=None):
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Emu

    prs = Presentation()
    if title:
        prs.core_properties.title = title
    prs.slide_width = Emu(SLIDE_W)
    prs.slide_height = Emu(SLIDE_H)
    blank = prs.slide_layouts[6]
    small = []
    for deck, no, path, w, h, tall in shots:
        slide = prs.slides.add_slide(blank)
        fill = slide.background.fill      # the paper the card sits on, so any
        fill.solid()                      # margin reads as deck, not as a gap
        fill.fore_color.rgb = RGBColor(*PAGE_BG)
        ratio = min(SLIDE_W / w, SLIDE_H / h)
        pw, ph = int(w * ratio), int(h * ratio)
        slide.shapes.add_picture(path, Emu((SLIDE_W - pw) // 2),
                                 Emu((SLIDE_H - ph) // 2), Emu(pw), Emu(ph))
        if tall:
            small.append((deck, no, int(h), round(pw / SLIDE_W * 100)))
    prs.save(out)
    return small


def contact_sheet(pptx_path, png_path, cols=5, cell_w=384):
    """Draw what a viewer will draw, reading only from inside the .pptx: the
    stored background colour, the embedded image, its stored geometry. This is
    not PowerPoint rendering the file - it is the file's own contents, painted."""
    from PIL import Image
    from pptx import Presentation

    prs = Presentation(pptx_path)
    sw, sh = prs.slide_width, prs.slide_height

    def render(slide, px_w):
        sc = px_w / sw
        rgb = str(slide.background.fill.fore_color.rgb)
        canvas = Image.new("RGB", (px_w, int(round(sh * sc))),
                           tuple(int(rgb[i:i + 2], 16) for i in (0, 2, 4)))
        pic = slide.shapes[0]
        img = Image.open(io.BytesIO(pic.image.blob)).convert("RGB")
        size = (max(1, int(round(pic.width * sc))),
                max(1, int(round(pic.height * sc))))
        canvas.paste(img.resize(size, Image.LANCZOS),
                     (int(round(pic.left * sc)), int(round(pic.top * sc))))
        return canvas

    slides = list(prs.slides)
    cell_h = int(round(cell_w * sh / sw))
    rows = (len(slides) + cols - 1) // cols
    sheet = Image.new("RGB",
                      (cols * (cell_w + 10) + 10, rows * (cell_h + 10) + 10),
                      (30, 40, 55))
    for i, s in enumerate(slides):
        sheet.paste(render(s, cell_w),
                    (10 + (i % cols) * (cell_w + 10),
                     10 + (i // cols) * (cell_h + 10)))
    sheet.save(png_path)
    return len(slides)


def main():
    ap = argparse.ArgumentParser(description="Export an HTML deck folder to .pptx")
    ap.add_argument("deck_folder")
    ap.add_argument("out", help="path of the .pptx to write")
    ap.add_argument("--port", type=int, default=8899)
    ap.add_argument("--decks", help="comma-separated file names, in slide order "
                                    "(default: every .html in the folder, sorted)")
    ap.add_argument("--frame", choices=("all", "fitted", "none"), default="all",
                    help="which slides keep the deck's card border and paper "
                         "(default: all)")
    ap.add_argument("--scale", type=int, default=2,
                    help="device pixel ratio (default 2; scaled-down slides "
                         "get one more)")
    ap.add_argument("--shots", help="where to keep the PNGs (default: a temp dir). "
                                    "Never point this inside the deck folder - the "
                                    "server would serve them and git would see them.")
    ap.add_argument("--title", help="presentation title (default: the deck's own "
                                    "<title>, when exporting a single deck)")
    ap.add_argument("--sheet", help="also write a contact sheet PNG of the result")
    a = ap.parse_args()

    folder = os.path.abspath(a.deck_folder)
    if a.decks:
        decks = [d.strip() for d in a.decks.split(",") if d.strip()]
    else:
        decks = sorted(os.path.basename(f)
                       for f in glob.glob(os.path.join(folder, "*.html")))
    if not decks:
        raise SystemExit("no .html decks in " + folder)

    base = "http://127.0.0.1:%d/" % a.port
    require_server(base, decks)

    shots_dir = a.shots or os.path.join(os.environ.get("TEMP", "/tmp"), "deck-shots")
    if os.path.abspath(shots_dir).startswith(folder + os.sep):
        raise SystemExit("--shots must not sit inside the deck folder")

    print("capturing from %s  (frame: %s)" % (base, a.frame))
    from playwright.sync_api import sync_playwright
    with sync_playwright() as p:
        shots = capture(p, base, decks, shots_dir, a.scale, a.frame)
    title = a.title or (deck_title(base, decks[0]) if len(decks) == 1 else None)
    small = build(shots, a.out, title)

    print("\nslides written : %d  (%s)" % (len(shots), ", ".join(decks)))
    print("output         : %s" % os.path.abspath(a.out))
    if title:
        print("title          : %s" % title)
    if a.sheet:
        n = contact_sheet(a.out, a.sheet)
        print("contact sheet  : %s  (%d slides)" % (os.path.abspath(a.sheet), n))
    if small:
        print("\nTALLER THAN 720px - fitted by height, so these sit smaller on the "
              "slide with deck paper either side (%d):" % len(small))
        for deck, no, h, pct in small:
            print("  %-14s %s  %dpx tall  -> uses %d%% of the slide width"
                  % (deck, no, h, pct))
        print("Under ~70% of the width the body text gets hard to read projected. "
              "Fix by trimming the slide to 720px in the HTML, or splitting it.")


if __name__ == "__main__":
    sys.exit(main())
